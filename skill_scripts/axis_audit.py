# -*- coding: utf-8 -*-
"""분기 자가진단 — 산출 직전에 축별 쏠림을 측정한다.

사람 눈은 쏠림에 관대하다. 실제로 WRTN 실행에서 '광고 각인'이 7/9 한 값에
쏠린 걸 못 보고 산출했다. 그래서 기계로 센다.

판정 기준 (표시되는 정수 %로 판정한다 — 아래 '반올림' 참조):
  ≥78%  🔴 도장  → 그 축은 규칙이 개인을 안 보고 있다. 고치기 전엔 산출 금지.
  56~77% 🟡 쏠림  → 임계값을 조정해 분기를 만든다.
                    단, 값이 코호트 보편(20대의 유튜브, 전 연령의 스마트폰)이면
                    임계 조정으로는 절대 안 갈라진다. 축의 층위를 한 단계 내려라
                    ('어느 채널' → '그 채널에서 무엇을').
  ≤55%  🟢 정상

반올림 — 이걸 틀려서 도구가 자기 존재 이유를 놓쳤다:
  7/9 = 77.78%다. 화면엔 '78%'로 찍히는데 `share >= 0.78`은 False라
  **도장이 아니라 쏠림으로 판정**됐다. 이 스크립트를 만들게 한 바로 그 WRTN
  사건이 통과해 버린 것이다. 5/9 = 55.56%('56%'로 표시) 역시 정상으로 샜다.
  세그당 3명·패널 9명이면 점유율이 11%씩 계단이라 항상 경계에서 터진다.
  → 그래서 **표시값과 판정값을 같게** 한다: pct = round(share*100)으로 판정.

세그별 감사 — 전체만 재면 원조 실패를 못 본다:
  이 스킬의 원조 실패는 "세그당 3명이 똑같음"(INF 도장)이다. 그런데
  세그A 전원 '상' / 세그B 전원 '중' / 세그C 전원 '하'면 패널 전체 최빈값은
  33%로 **정상 통과**한다 — 세그 간 차이가 세그 내 도장을 완벽히 가린다.
  → '세그' 컬럼이 있으면 전체 감사에 더해 **세그별로도** 잰다.

예외(도장이 아님):
  1. 의도적 고정값 (브랜드 선호 = 미부여)
  2. 세그 정의를 되풀이하는 동어반복 축 → 애초에 카드에 올리지 말 것
  3. **사용자가 게이트 B에서 정량 예고를 받고 감수하기로 확정한 쏠림**
     (예: 단일 연령대 선택 → 인지 경계 쏠림). 도장이 아니라 '기록된 트레이드오프'다.
     `--accepted <축이름>`으로 표시하면 ⚪로 분류하고, 산출물에 수치와 함께 명시한다.

사용:
  python3 axis_audit.py <panel_carriers.xlsx>
  python3 axis_audit.py <panel_carriers.xlsx> --ignore 근거 --accepted "AI 인지 경계"
"""
import argparse, re, sys
from collections import Counter, defaultdict

# 값 정규화 — 근거만 떼고 값은 건드리지 않는다.
#
# 예전엔 `v.split("(")[0]`로 괄호를 통째로 잘랐다. 근거가 값 뒤에 괄호로
# 붙던 시절('유튜브 (개방 t60)')엔 맞았지만, 지금은 **값 자체가 괄호를 쓴다**
# ('유튜브(운동) · 알고리즘 추천 수동 소비'). 그래서 서로 다른 9종이 전부
# '유튜브'로 뭉개져 **감사기가 89% 가짜 도장을 만들어냈다** — 이미 고쳐놓은
# 축을 도장이라 우기며 산출을 막았다. 측정 도구가 측정 대상을 오염시킨 것이다.
#
# 그러니 근거로 확인되는 것만 뗀다: (a) '←' 뒤 전부, (b) **말미** 괄호이면서
# 안에 근거 표지가 든 것. 표지가 없는 괄호는 값의 일부다 — 그냥 둔다.
#   근거다:   '상 (개방성 매우 높음, "게임")' · '유튜브 (개방 t60)' · '중 (INF·연령 파생)'
#   값이다:   '유튜브(운동) · 인스타그램' · '들어는 봤는데 뭔지 모름 (접점 3개)'
# 완벽히 가를 수는 없다(자연어다). 그래서 근거 표지를 **열거**해 보수적으로 뗀다 —
# 못 떼면 고유값이 늘어 도장을 놓칠 뿐이지만, 잘못 떼면 없는 도장을 만들어
# 멀쩡한 축을 반려시킨다. 후자가 훨씬 해롭다.
# 가장 안전한 길은 애초에 근거를 값에 섞지 않는 것이다: '값 ← 근거' 또는 별도 '_근거' 컬럼.
_TAGS = r"Q|Q-derived|Q\+INF|T|T:[^\]]+|INF|INF·가설|가설|경계:모름"
_EV_MARK = (r"\[(?:" + _TAGS + r")\]"          # provenance 태그
            r"|\bt\d{2}\b"                      # t점수
            r"|개방성|성실성|외향성|우호성|신경성"   # Big5 라벨
            r"|파생|근거|신호|프록시")            # 근거 서술 표지
_EV_PAREN = re.compile(r"\s*\((?=[^()]*(?:" + _EV_MARK + r"))[^()]*\)\s*$")
_EV_TAG = re.compile(r"\s*\[(?:" + _TAGS + r")\]\s*$")


def norm(v):
    v = str(v).split("←")[0]
    prev = None
    while v != prev:                       # 근거 괄호·태그가 겹쳐 붙는 경우
        prev = v
        v = _EV_PAREN.sub("", v)
        v = _EV_TAG.sub("", v)
    return v.strip()

# 식별자·원본 데모 — 감사 대상이 아니다(합성 축이 아니라 Q 필드다).
# 정확 일치로만 스킵한다. 부분문자열로 스킵하면 '가구 구매력(소득 프록시)' 같은
# **합성 축**이 "소득" 포함이라는 이유로 감사에서 빠져나간다.
SKIP_EXACT = {"세그", "stance", "uuid", "이름", "나이", "성별", "직업", "소득",
              "지역", "시군구", "선발점수", "세그근거", "가구", "주거", "학력",
              "전공", "경제활동"}
# 설계상 전원 동일·고정값이거나 축이 아닌 행 — 도장 대상이 아니다(SKILL '예외' 1·상식 기저).
# 헤더에 괄호 설명이 붙어도("브랜드 선호(비렉스·코웨이)") 빠지게 앞부분으로 맞춘다.
SKIP_PREFIX = ("Big5", "상식 기저", "브랜드 선호", "기본 프로필", "FGD 화법")
SEG_KEYS = ("세그", "stance")


def judge(cnt, n, accepted=False):
    """표시되는 정수 %로 판정한다 — 표시와 판정이 어긋나면 도구를 못 믿는다."""
    pct = round(cnt / n * 100)
    if accepted:
        return "accepted", pct
    if pct >= 78:
        return "stamp", pct
    if pct >= 56:
        return "tilt", pct
    return "ok", pct


def audit(hdr, rows, skip, accepted, label=""):
    n = len(rows)
    buckets = defaultdict(list)
    for i, h in enumerate(hdr):
        if h in skip:
            continue
        if any(s in h for s in skip if s.endswith("_근거")):
            continue
        # 근거가 달라도 값이 같으면 쏠림이다 → 근거만 떼고 값을 본다 (norm 참조)
        vals = [norm(r[i]) for r in rows]
        if len({v for v in vals if v and v != "None"}) == 0:
            continue
        c = Counter(vals)
        top, cnt = c.most_common(1)[0]
        kind, pct = judge(cnt, n, accepted=(h in accepted))
        buckets[kind].append(
            f"  {h:24} 최빈 '{top[:22]}' {cnt}/{n} ({pct}%) · 고유값 {len(c)}")
    return buckets, n


def stamped_axes(buckets):
    """도장 난 축 이름만 뽑는다 (세그 간 반복 여부를 보려고)."""
    return {l.strip().split("  ")[0].strip() for l in buckets.get("stamp", [])}


def show(buckets, n, title):
    order = [("stamp", "🔴 도장 — 고치기 전엔 산출하지 말 것"),
             ("tilt", "🟡 쏠림 — 임계값 조정, 또는 값이 코호트 보편이면 축의 층위를 내릴 것"),
             ("accepted", "⚪ 감수 확정 — 사용자가 예고받고 수용한 트레이드오프 (산출물에 명시할 것)"),
             ("ok", "🟢 정상")]
    print(f"\n{'='*68}\n{title} (n={n})\n{'='*68}")
    for k, t in order:
        if buckets.get(k):
            print(t)
            for l in buckets[k]:
                print(l)
            print()
    return len(buckets.get("stamp", []))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--ignore", nargs="*", default=[],
                    help="감사에서 뺄 컬럼(정확 일치). 근거 컬럼은 자동 제외.")
    ap.add_argument("--accepted", nargs="*", default=[],
                    help="게이트 B에서 사용자가 예고받고 감수한 축 → 도장 아님(예외 3번)")
    a = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        print("openpyxl 필요"); sys.exit(1)
    ws = openpyxl.load_workbook(a.xlsx).active
    hdr = [str(c.value) for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        print("빈 패널"); return

    skip = SKIP_EXACT | set(a.ignore) | {h for h in hdr if h.endswith("_근거") or h.startswith(SKIP_PREFIX)}
    accepted = set(a.accepted)

    b, n = audit(hdr, rows, skip, accepted)
    bad = show(b, n, "패널 전체")

    # ── 세그별 — 전체 최빈값이 세그 내 도장을 가린다 ──
    segi = next((i for i, h in enumerate(hdr) if h in SEG_KEYS), None)
    seg_bad = 0
    hits = Counter()      # 축 → 도장 난 세그 수
    nsegs = 0
    if segi is not None:
        groups = defaultdict(list)
        for r in rows:
            groups[r[segi]].append(r)
        print(f"\n{'#'*68}\n세그별 감사 — 세그 안에서 갈라지지 않으면 라벨 3장이다\n{'#'*68}")
        for seg, g in groups.items():
            if len(g) < 2:
                continue
            nsegs += 1
            gb, gn = audit(hdr, g, skip, accepted)
            seg_bad += show(gb, gn, f"[{seg}]")
            for ax in stamped_axes(gb):
                hits[ax] += 1
    else:
        print("\n(‘세그’ 컬럼이 없어 세그별 감사를 건너뜀 — 세그 내 도장은 못 잡는다)")

    print(f"\n{'='*68}")
    print(f"판정: 전체 도장 {bad}건 · 세그 내 도장 {seg_bad}건")

    # 우연과 고장을 가른다 — n=3에서 3수준 축이 3/3으로 뭉칠 확률은 약 11%다.
    # 축 13개 × 세그 3개면 우연만으로도 4건쯤 뜬다. 그러니 건수를 세지 말고
    # **같은 축이 여러 세그에서 반복 도장인지**를 봐라. 반복은 우연이 아니라
    # 규칙이 개인을 안 보고 있다는 뜻이다.
    rep, one = {}, {}
    if hits:
        rep = {a: c for a, c in hits.items() if c >= 2}
        one = {a: c for a, c in hits.items() if c == 1}
        if rep:
            print(f"\n🔴 반복 도장 — 규칙이 개인을 안 본다 (여러 세그에서 동시에 뭉침)")
            for a, c in sorted(rep.items(), key=lambda x: -x[1]):
                print(f"   {a} — {c}/{nsegs}개 세그에서 도장. 이 축의 파생 규칙을 고쳐라.")
        if one:
            print(f"\n🟡 단일 세그 도장 — 우연일 수 있다 (n이 작으면 3/3은 흔하다)")
            for a in one:
                print(f"   {a} — 1개 세그. 세그 정의를 되풀이하는 축이면 카드에서 빼라(동어반복).")

    if bad or rep:
        print("\n→ 고치기 전엔 산출하지 마라.")
        print("   값이 코호트 보편이면 임계 조정이 아니라 축의 층위를 내려라.")
        print("   의도적 고정·동어반복 축이면 --ignore, 감수 확정이면 --accepted.")
        sys.exit(1)          # 파이프라인에서 산출을 막으려면 종료코드가 필요하다
    elif one:
        print("\n→ 🟡 반복 도장은 없다. 단일 세그 도장이 동어반복인지만 확인하고 진행.")
    else:
        print("→ 🟢 산출 진행 가능")


if __name__ == "__main__":
    main()
