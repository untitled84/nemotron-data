# -*- coding: utf-8 -*-
"""패널 산출 렌더러: 세그먼트 카드 HTML + 캐리어 xlsx + FGD용 panel.json.

render_panel(panel, html_path)        # 사람이 보는 카드 (아래 contract)
save_xlsx(records, xlsx_path)         # 검증·핸드오프용 표 (records: list[dict])
save_panel_json(panel, json_path)     # persona-qual-fgd / nemotron-fgd-studio.html 입력

panel = {
  "title": str, "subtitle": str, "note": str,
  "meta": {"brand": str, "category": str, ...},        # JSON에만 실린다(HTML은 무시)
  "segments": [
    {"name": str, "color": "#hex", "criteria": str,
     "people": [
        {"header": "이름 34여 서울 직업", "uuid": "....",
         "rows": [ ("라벨","값") or ("라벨","값","Q"|"Q-derived"|"T"|"INF"|"가설") ],
         # ↓ JSON 전용(선택). render_panel은 읽지 않으므로 HTML 산출은 그대로다.
         "name": "홍길동",                # 없으면 header 첫 토큰
         "contradiction": "…",            # 모순 (FGD에서 흔들릴 지점)
         "boundary": "…",                 # 인지 경계 [경계:모름]
         "raw": { …니모트론 원본 51컬럼 전부… }   # big5는 아래 규칙대로 접힌다
        }
     ]}
  ]}
provenance 태그는 각 값의 근거를 드러내려는 것: Q=원본필드, Q-derived=카드신호 연역,
T=외부통계/웹, INF=성향추론(합성), 가설=압박테스트 대상(브랜드 선호 등).
"""
import html as _h
import json as _j

# provenance.md의 태그 체계를 온전히 따른다. 미등록 태그는 폴백(INF 보라)으로 찍혀
# 근거 두께를 오독하게 만든다 — '경계:모름'은 INF가 아니라 이 스킬의 핵심 산출이다.
_TAGCLS={"Q":"q","Q-derived":"q","T":"t","T:웹-레퍼토리":"t","INF":"i","가설":"h","경계:모름":"b"}

def render_panel(panel, html_path):
    segs=panel["segments"]
    def tag(t):
        cls=_TAGCLS.get(t,"i")
        return f'<span class="tg {cls}">{_h.escape(t)}</span>'
    body=""
    for s in segs:
        cards=""
        for p in s["people"]:
            rows=""
            for r in p["rows"]:
                lab,val=r[0],r[1]; tg=(" "+tag(r[2])) if len(r)>2 else ""
                rows+=f'<div class="li"><b>{_h.escape(lab)}</b> {_h.escape(str(val))}{tg}</div>'
            uu=f'<div class="uu">uuid {_h.escape(str(p.get("uuid","")))[:12]}…</div>' if p.get("uuid") else ""
            cards+=f'<div class="c"><div class="hd">{_h.escape(p["header"])}</div>{rows}{uu}</div>'
        body+=(f'<section><div class="band" style="background:{s.get("color","#444")}">'
               f'<h2>{_h.escape(s["name"])}</h2><span class="cr">{_h.escape(s.get("criteria",""))}</span></div>'
               f'<div class="grid">{cards}</div></section>')
    doc=f'''<!DOCTYPE html><html lang="ko"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{_h.escape(panel.get("title","패널"))}</title>
<style>:root{{--bg:#f5f6f8;--card:#fff;--ink:#1c2230;--sub:#5b6472;--line:#e4e7ec}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font-family:-apple-system,BlinkMacSystemFont,"Apple SD Gothic Neo","Malgun Gothic",sans-serif;line-height:1.5;padding:26px 20px 60px}}
.wrap{{max-width:1180px;margin:0 auto}}h1{{font-size:23px;margin:0 0 5px}}.sub{{color:var(--sub);font-size:13.5px;margin:0 0 14px;max-width:920px}}
.note{{background:#eef7f1;border:1px solid #cfe6d8;border-radius:10px;padding:11px 15px;font-size:12.5px;color:#28553f;margin-bottom:20px}}
.tg{{display:inline-block;font-size:10px;padding:0 6px;border-radius:10px;margin-left:3px;vertical-align:middle}}
.tg.q{{background:#e6f4ee;color:#1f6f5c}}.tg.t{{background:#e9f0fb;color:#2b5d9e}}.tg.i{{background:#f3ecfb;color:#5b3fa0}}.tg.h{{background:#fbeede;color:#9a5a12}}.tg.b{{background:#eceef1;color:#4a5560}}
section{{margin-bottom:24px}}.band{{display:flex;align-items:baseline;gap:10px;padding:8px 13px;border-radius:9px;color:#fff;margin-bottom:12px}}
.band h2{{margin:0;font-size:16px}}.band .cr{{font-size:12px;opacity:.93}}
.grid{{display:grid;grid-template-columns:1fr 1fr 1fr;gap:13px}}@media(max-width:820px){{.grid{{grid-template-columns:1fr}}}}
.c{{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:13px 14px}}
.hd{{font-size:14px;font-weight:700;margin-bottom:7px}}.li{{font-size:12px;margin:5px 0}}.li b{{color:var(--sub);font-weight:700;font-size:11px}}
.uu{{font-size:10px;color:#9aa2ad;margin-top:7px;font-family:monospace}}</style></head><body><div class="wrap">
<h1>{_h.escape(panel.get("title","패널"))}</h1><p class="sub">{_h.escape(panel.get("subtitle",""))}</p>
<div class="note">{panel.get("note","")}</div>{body}</div></body></html>'''
    with open(html_path,"w",encoding="utf-8") as f: f.write(doc)
    return html_path

B5_KEYS = [("openness","개방성"),("conscientiousness","성실성"),("extraversion","외향성"),
           ("agreeableness","우호성"),("neuroticism","신경성")]


def _fold_big5(raw):
    """Big5 5컬럼을 raw.big5 한 덩이로 접는다.

    니모트론의 Big5는 JSON **문자열**이다: '{"t_score":31,"label":"매우 낮음","description":"…"}'
    파싱하지 않고 그대로 넘기면 FGD 프롬프트에 원시 JSON이 통째로 박혀 소음이 되고,
    소비 측(nemotron-fgd-studio)은 `Object.entries(raw.big5)`로 k:v를 뽑으므로
    **평탄한 dict**를 기대한다. → 't66 · 높음' 형태로 접어 사람도 모델도 읽게 한다.
    """
    out = {}
    for k, ko in B5_KEYS:
        v = raw.pop(k, None)
        if isinstance(v, str):
            try: v = _j.loads(v)
            except Exception: pass
        if isinstance(v, dict) and v.get("t_score") is not None:
            lab = v.get("label") or ""
            out[ko] = f"t{int(v['t_score'])}" + (f" · {lab}" if lab else "")
        elif v is not None:
            out[ko] = v
    return out


def save_panel_json(panel, json_path):
    """FGD 입력용 panel.json — 니모트론 원본 전부 + 합성 축 전부.

    왜 별도 산출인가: HTML 카드는 **사람이 읽는 증류본**이라 토픽에 필요한 축만 올린다.
    그런데 FGD에서 페르소나가 답할 땐 카드에 안 올린 배경(직업 서사·돈 태도·건강)까지
    근거가 되어야 그 사람처럼 말한다. 카드를 소음으로 채우지 않으면서 근거는 다 넘기려면
    **보여주는 것(HTML)과 넘기는 것(JSON)을 분리**하는 수밖에 없다.

    소비 측 계약(nemotron-fgd-studio.html · loadPanel/normPerson/nemoRoster에서 확인):
      {meta, segments:[{name,color,criteria,personas:[
         {uuid, header, name, cell, rows:[[라벨,값,태그]], contradiction, boundary, raw:{…, big5:{}}}]}]}
      - rows의 태그는 세 번째 원소. 2원소(태그 없음)도 허용된다.
      - raw는 니모트론 원본 51컬럼 전부. 스튜디오는 그중 일부만 프롬프트에 찍지만,
        전부 실어야 다른 토픽의 FGD에서도 근거로 쓸 수 있다(카드 하나로 여러 FGD).
    """
    segs = []
    missing_raw = []
    for s in panel["segments"]:
        people = []
        for p in s["people"]:
            raw = dict(p.get("raw") or {})
            if not raw:
                missing_raw.append(p.get("header", "?"))
            big5 = _fold_big5(raw)
            if big5:
                raw["big5"] = big5
            raw = {k: v for k, v in raw.items() if v is not None and v != ""}
            people.append({
                "uuid": str(p.get("uuid", "")),
                "header": p.get("header", ""),
                "name": p.get("name") or (p.get("header", "").split(" ")[0] or "참여자"),
                "cell": s["name"],
                "rows": [list(r) for r in p.get("rows", [])],
                "contradiction": p.get("contradiction", "") or "",
                "boundary": p.get("boundary", "") or "",
                "raw": raw,
            })
        segs.append({"name": s["name"], "color": s.get("color", "#15a89a"),
                     "criteria": s.get("criteria", ""), "personas": people})
    meta = dict(panel.get("meta") or {})
    meta.setdefault("title", panel.get("title", ""))
    meta.setdefault("subtitle", panel.get("subtitle", ""))
    meta.setdefault("note", panel.get("note", ""))
    doc = {"meta": meta, "segments": segs}
    with open(json_path, "w", encoding="utf-8") as f:
        _j.dump(doc, f, ensure_ascii=False, indent=1)
    n = sum(len(s["personas"]) for s in segs)
    if missing_raw:
        print(f"  ⚠ raw 없는 인물 {len(missing_raw)}명 — FGD에서 배경 근거 없이 답하게 된다: "
              + ", ".join(missing_raw[:3]))
    return json_path


def save_xlsx(records, xlsx_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    cols=[]
    for r in records:
        for k in r:
            if k not in cols: cols.append(k)
    wb=Workbook(); ws=wb.active; ws.title="panel"
    ws.append(cols)
    for c in range(1,len(cols)+1):
        cell=ws.cell(1,c); cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="2b5d9e")
    for r in records:
        ws.append([ ("" if r.get(k) is None else str(r.get(k))) for k in cols ])
    for i,k in enumerate(cols,1):
        w=min(60,max(10,max([len(str(k))]+[len(str(r.get(k,""))) for r in records[:200]])+2))
        ws.column_dimensions[ws.cell(1,i).column_letter].width=w
    ws.freeze_panes="A2"
    wb.save(xlsx_path); return xlsx_path
